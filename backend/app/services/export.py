"""PDF / Excel export for shift tables."""

from __future__ import annotations

import calendar
import io
from collections import defaultdict
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.models.employee import Employee
from app.models.shift import ShiftAssignment
from app.services.holidays import is_holiday


def _register_jp_font() -> str:
    try:
        pdfmetrics.registerFont(UnicodeCIDFont("HeiseiKakuGo-W5"))
        return "HeiseiKakuGo-W5"
    except Exception:
        return "Helvetica"


def _build_grid(
    year: int,
    month: int,
    assignments: list[ShiftAssignment],
    employees: list[Employee],
) -> tuple[list[date], list[list[str]]]:
    _, ndays = calendar.monthrange(year, month)
    days = [date(year, month, d) for d in range(1, ndays + 1)]

    per_emp_day: dict[int, dict[date, list[str]]] = defaultdict(lambda: defaultdict(list))
    for a in assignments:
        # 実時間帯で表示する（例 20:00-01:00）。深夜跨ぎもそのまま。
        cell = f"{a.start_time.strftime('%H:%M')}-{a.end_time.strftime('%H:%M')}"
        per_emp_day[a.employee_id][a.target_date].append(cell)

    rows: list[list[str]] = []
    header = ["従業員"] + [f"{d.day}({'月火水木金土日'[d.weekday()]})" for d in days]
    rows.append(header)
    for e in employees:
        row = [e.name]
        for d in days:
            row.append("/".join(per_emp_day[e.id].get(d, [])) or "-")
        rows.append(row)
    return days, rows


def export_pdf(
    year: int,
    month: int,
    assignments: list[ShiftAssignment],
    employees: list[Employee],
) -> bytes:
    font_name = _register_jp_font()
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), title=f"Shift-{year}-{month:02d}")
    styles = getSampleStyleSheet()
    story = []
    title = Paragraph(
        f"<font name='{font_name}' size='16'>{year}年{month}月 シフト表</font>",
        styles["Normal"],
    )
    story.append(title)
    story.append(Spacer(1, 12))

    days, rows = _build_grid(year, month, assignments, employees)
    table = Table(rows, repeatRows=1)
    style = TableStyle(
        [
            ("FONTNAME", (0, 0), (-1, -1), font_name),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
            ("ALIGN", (1, 1), (-1, -1), "CENTER"),
        ]
    )
    for i, d in enumerate(days, start=1):
        if d.weekday() == 5:
            style.add("TEXTCOLOR", (i, 0), (i, 0), colors.blue)
        elif d.weekday() == 6 or is_holiday(d):
            style.add("TEXTCOLOR", (i, 0), (i, 0), colors.red)

    table.setStyle(style)
    story.append(table)
    doc.build(story)
    return buffer.getvalue()


def export_excel(
    year: int,
    month: int,
    assignments: list[ShiftAssignment],
    employees: list[Employee],
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = f"{year}-{month:02d}"

    days, rows = _build_grid(year, month, assignments, employees)

    header_fill = PatternFill(start_color="EEEEEE", end_color="EEEEEE", fill_type="solid")
    sat_fill = PatternFill(start_color="D6EAF8", end_color="D6EAF8", fill_type="solid")
    hol_fill = PatternFill(start_color="FADBD8", end_color="FADBD8", fill_type="solid")

    for ci, val in enumerate(rows[0], start=1):
        cell = ws.cell(row=1, column=ci, value=val)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        if ci >= 2:
            d = days[ci - 2]
            if d.weekday() == 5:
                cell.fill = sat_fill
            elif d.weekday() == 6 or is_holiday(d):
                cell.fill = hol_fill

    for ri, row in enumerate(rows[1:], start=2):
        for ci, val in enumerate(row, start=1):
            cell = ws.cell(row=ri, column=ci, value=val)
            if ci >= 2:
                cell.alignment = Alignment(horizontal="center")

    ws.column_dimensions["A"].width = 14

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
